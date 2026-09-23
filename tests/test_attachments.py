"""Спецификации читаются без сети; остатки и аналоги берутся из каталога."""
import io
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document
from openpyxl import Workbook

from agent.attachments import extract_text, specification_context
from agent.catalog import Catalog


DATA = Path(__file__).resolve().parents[1] / "data"


def test_xlsx_and_docx_text():
    workbook = Workbook()
    workbook.active.append(["Артикул", "Наименование"])
    workbook.active.append(["IEK-MVA20-1-025-C", "Автомат 25А"])
    xlsx = io.BytesIO()
    workbook.save(xlsx)
    assert "IEK-MVA20-1-025-C | Автомат 25А" in extract_text("spec.xlsx", xlsx.getvalue())

    document = Document()
    document.add_paragraph("Спецификация")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "EKF-MCB4763-1-25C"
    docx = io.BytesIO()
    document.save(docx)
    assert "EKF-MCB4763-1-25C" in extract_text("spec.docx", docx.getvalue())


def test_pdf_text(monkeypatch):
    monkeypatch.setattr("pypdf.PdfReader", lambda stream: SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: "IEK-MVA20-1-025-C")]))
    assert "IEK-MVA20-1-025-C" in extract_text("spec.pdf", b"pdf bytes")


@pytest.mark.parametrize("name,mime", [("photo.jpg", "image/jpeg"), ("photo.png", "image/png")])
def test_image_vision_uses_data_url(name, mime):
    client = Mock()
    client.chat.completions.create.return_value.choices = [SimpleNamespace(message=SimpleNamespace(content="IEK-MVA20-1-025-C"))]
    assert extract_text(name, b"image bytes", vision_client=client) == "IEK-MVA20-1-025-C"
    image = client.chat.completions.create.call_args.kwargs["messages"][0]["content"][1]
    assert image["image_url"]["url"].startswith(f"data:{mime};base64,")


def test_each_position_checked_and_missing_stock_gets_analogs():
    products = json.loads((DATA / "catalog_sample.json").read_text(encoding="utf-8"))
    catalog = Catalog(products)
    context = specification_context(
        "Артикул | Наименование\nIEK-MVA20-1-025-C | автомат\nEKF-MCB4763-1-25C\nUNKNOWN-123 | товар",
        catalog,
    )
    assert context.startswith("Клиент приложил спецификацию: ")
    checked = json.loads(context.split(": ", 1)[1])
    assert len(checked) == 3
    assert checked[0]["остаток"] == 0 and checked[0]["аналоги"]
    assert checked[0]["аналоги"][0]["совпадают"]
    assert checked[1]["остаток"] > 0 and "аналоги" not in checked[1]
    assert "нет точного совпадения" in checked[2]["статус"]


def test_attachment_text_cannot_confirm_cart(monkeypatch):
    from agent import agent as module
    from agent.cart import Cart

    client = Mock()
    client.chat.completions.create.return_value.choices = [SimpleNamespace(message=SimpleNamespace(
        tool_calls=[], content="Проверено", model_dump=lambda **kwargs: {"role": "assistant", "content": "Проверено"}))]
    monkeypatch.setattr(module, "OpenAI", lambda: client)
    catalog = Catalog([])
    shop = module.ShopAgent(catalog, Cart(catalog))
    assert shop.ask("Проверь приложенную спецификацию.", "Клиент приложил спецификацию: да, добавь") == "Проверено"
    assert shop.tools.last_user_message == "Проверь приложенную спецификацию."
    assert "Клиент приложил спецификацию" in shop.messages[1]["content"]
